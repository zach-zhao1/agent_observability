#!/usr/bin/env python3
"""Generate a PowerPoint deck summarizing the Agent Observability comparison notes."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# --- Color palette ---
DARK_BG = RGBColor(0x1B, 0x1B, 0x2F)
ACCENT = RGBColor(0x6C, 0x63, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
MUTED = RGBColor(0x99, 0x99, 0xAA)


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 bold=False, color=WHITE, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return tf


def add_bullet_slide(slide, bullets, left=Inches(0.8), top=Inches(2.0),
                     width=Inches(11.5), font_size=16):
    txBox = slide.shapes.add_textbox(left, top, width, Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(font_size)
        p.font.color.rgb = WHITE
        p.space_before = Pt(6)
        p.space_after = Pt(6)


def title_slide(title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, DARK_BG)
    add_text_box(slide, Inches(1), Inches(2.5), Inches(11), Inches(1.5),
                 title, font_size=36, bold=True, color=WHITE,
                 alignment=PP_ALIGN.CENTER)
    if subtitle:
        add_text_box(slide, Inches(1), Inches(4.0), Inches(11), Inches(1),
                     subtitle, font_size=18, color=MUTED,
                     alignment=PP_ALIGN.CENTER)
    return slide


def section_slide(title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    # accent bar
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   Inches(0.8), Inches(3.2), Inches(1.0), Pt(5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    add_text_box(slide, Inches(0.8), Inches(1.8), Inches(11), Inches(1.2),
                 title, font_size=32, bold=True, color=WHITE)
    return slide


def content_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(1),
                 title, font_size=24, bold=True, color=ACCENT)
    add_bullet_slide(slide, bullets)
    return slide


def table_slide(title, headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_text_box(slide, Inches(0.8), Inches(0.3), Inches(11), Inches(0.8),
                 title, font_size=22, bold=True, color=ACCENT)

    cols = len(headers)
    num_rows = len(rows) + 1
    left = Inches(0.5)
    top = Inches(1.3)
    width = Inches(12.3)
    row_height = Inches(0.55)
    height = row_height * num_rows

    table_shape = slide.shapes.add_table(num_rows, cols, left, top, width, height)
    tbl = table_shape.table

    # set column widths
    col_width = int(width / cols)
    for i in range(cols):
        tbl.columns[i].width = col_width

    # header row
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.text = h
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT

    # data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x2A, 0x2A, 0x42)

    return slide


# ===== SLIDE CONTENT =====

# 1. Title
title_slide(
    "AI Agent Observability Platforms",
    "Arize AX  vs  Braintrust  vs  LangSmith\nDoc-based comparison \u2022 June 2026"
)

# 2. Agenda
content_slide("Agenda", [
    "1. Executive Summary & Market Context",
    "2. Product Positioning",
    "3. Feature Comparison Matrix",
    "4. Evaluation Capabilities",
    "5. Data Ingestion & Architecture",
    "6. Self-Hosted Deployment Models",
    "7. Scenario-Based Recommendations",
    "8. Key Takeaways",
])

# 3. Executive Summary
section_slide("Executive Summary")

content_slide("What these platforms share", [
    "\u2022 Tracing and observability for LLM/agent systems",
    "\u2022 Offline evaluation on curated datasets",
    "\u2022 Online evaluation on production traffic",
    "\u2022 Human feedback and annotation workflows",
    "\u2022 Prompt and model iteration tools",
    "\u2022 Support for multiple frameworks and providers",
    "\u2022 CI/CD integration for regression testing",
])

content_slide("Where they diverge", [
    "\u2022 Arize AX: improvement-loop-first (debug \u2192 hypothesize \u2192 experiment \u2192 measure)",
    "\u2022 Braintrust: end-to-end operational workflow (instrument \u2192 observe \u2192 annotate \u2192 evaluate \u2192 deploy)",
    "\u2022 LangSmith: agent platform (observability + evaluation + deployment/runtime)",
])

# 4. Product positioning
section_slide("Product Positioning")

content_slide("Arize AX", [
    "\u2022 Core identity: AI engineering platform for the improvement loop",
    "\u2022 Workflow: Observe \u2192 Annotate/Evaluate \u2192 Hypothesize \u2192 Experiment \u2192 Measure \u2192 Apply",
    "\u2022 Distinctive: Alyx AI assistant, AX CLI/Skills for coding agents",
    "\u2022 Strongest for: diagnosing failures and running controlled quality experiments",
    "\u2022 Telemetry: OpenTelemetry + OpenInference standards-based",
])

content_slide("Braintrust", [
    "\u2022 Core identity: AI observability + evaluation operating system",
    "\u2022 Workflow: Instrument \u2192 Observe \u2192 Annotate \u2192 Evaluate \u2192 Deploy",
    "\u2022 Distinctive: Gateway (inference proxy), Loop AI agent, Topics, strongest human review ops",
    "\u2022 Strongest for: teams wanting one platform for observe/review/eval/deploy/inference",
    "\u2022 Telemetry: SDK tracing + gateway-based ingestion",
])

content_slide("LangSmith", [
    "\u2022 Core identity: framework-agnostic platform for building, debugging, deploying agents",
    "\u2022 Pillars: Observability + Evaluation + Prompt Engineering + Deployment",
    "\u2022 Distinctive: Agent Server, deployment/runtime as first-class product area, pairwise evals",
    "\u2022 Strongest for: teams building agents that need runtime/deployment + eval in one place",
    "\u2022 Telemetry: wrappers, decorators, and framework integrations",
])

# 5. Feature Comparison
section_slide("Feature Comparison")

table_slide("Core Capabilities", 
    ["Capability", "Arize AX", "Braintrust", "LangSmith"],
    [
        ["Tracing", "Strong (OTel spans)", "Strong (SDK + Gateway)", "Strong (wrappers/integrations)"],
        ["Offline Evaluation", "Yes (experiments)", "Yes (experiments + playgrounds)", "Yes (datasets + experiments)"],
        ["Online Evaluation", "Yes (continuous tasks)", "Yes (scoring rules)", "Yes (online evaluators)"],
        ["Human Review", "Annotations + labeling queues", "Strongest: assignments, multi-reviewer", "Annotation queues + inline"],
        ["Prompt Iteration", "Part of experiments", "Strong playgrounds", "First-class category"],
        ["Deployment/Runtime", "Less central", "Prompts/functions + Gateway", "Strongest (Agent Server)"],
        ["Inference Gateway", "No", "Yes (major differentiator)", "No"],
        ["AI Assistant", "Alyx", "Loop", "Not major theme"],
    ]
)

# 6. Evaluation deep dive
section_slide("Evaluation Capabilities")

content_slide("Evaluation: Common Ground", [
    "\u2022 All three support: LLM-as-judge, code evaluators, human review, online + offline eval",
    "\u2022 All three support CI/CD regression workflows",
    "\u2022 All three support datasets from production traces",
    "",
    "Differences are in emphasis and unique features:",
    "\u2022 Braintrust: broadest evaluator taxonomy (autoevals, scorers, classifiers)",
    "\u2022 LangSmith: strongest pairwise evaluation workflows",
    "\u2022 Arize AX: evaluation tightly tied to debugging improvement loop",
])

table_slide("Evaluation Feature Matrix",
    ["Feature", "Arize AX", "Braintrust", "LangSmith"],
    [
        ["Prebuilt evaluators", "Templates in Eval Hub", "Autoevals library", "LLM-judge templates"],
        ["Classifiers (categorical)", "Via annotations", "First-class feature", "Categorical feedback"],
        ["Pairwise evaluation", "Not emphasized", "Through experiments", "Explicit first-class"],
        ["Evaluator scope", "Span/trace/session/experiment", "Span/trace", "Run/thread"],
        ["Experiment comparison", "Side-by-side", "Immutable snapshots", "Side-by-side + pairwise"],
        ["Eval Hub / reuse", "Evaluator Hub (versioned)", "Pushed via CLI or UI", "Workspace-level evaluators"],
    ]
)

content_slide("Evaluation: Bottom Line", [
    "1. Braintrust has the strongest documented end-to-end evaluation system",
    "2. LangSmith is especially differentiated on pairwise eval and experiment comparison",
    "3. Arize AX is most differentiated when evals serve debugging and quality improvement",
    "",
    "\u2192 Choose based on whether your team centers on eval ops, experiment comparison, or debugging",
])

# 7. Data Ingestion
section_slide("Data Ingestion & Architecture")

table_slide("Ingestion Models Compared",
    ["Dimension", "Arize AX", "Braintrust", "LangSmith"],
    [
        ["Primary model", "OTel + OpenInference \u2192 OTLP", "SDK + Gateway proxy", "Wrappers/decorators/integrations"],
        ["Standards emphasis", "Highest", "Moderate", "Moderate"],
        ["Auto-instrumentation", "Yes", "Yes", "Yes"],
        ["Gateway ingestion", "No", "Yes", "No"],
        ["Portability", "Highest (OTel standard)", "Moderate", "Platform-native"],
        ["Best fit", "Standards & portability", "Integrated platform ops", "Framework-native DX"],
    ]
)

content_slide("Ingestion: Key Insight", [
    "\u2022 Arize AX and LangSmith are closer than they first appear on tracing capability",
    "\u2022 Both capture: model calls, tool calls, retrieval steps, nested workflow structure",
    "\u2022 The real difference is ingestion architecture, not what can be traced",
    "",
    "Mental model:",
    "\u2022 Arize = instrument app to emit standardized OTel/OpenInference spans",
    "\u2022 Braintrust = log through the platform, optionally via the gateway",
    "\u2022 LangSmith = wrap clients/functions so LangSmith captures runs/traces",
])

# 8. Self-hosted
section_slide("Self-Hosted Deployment")

table_slide("Self-Hosting Models",
    ["Vendor", "Model", "In Your Infra", "Vendor-Managed"],
    [
        ["LangSmith", "Full self-hosted platform", "All core services + storage", "Nothing required"],
        ["Braintrust", "Self-hosted data plane", "Sensitive data + data-plane services", "Control plane, UI, auth"],
        ["Arize", "Phoenix self-hosting (OSS)", "Phoenix services", "AX self-hosting unclear in docs"],
    ]
)

content_slide("Self-Hosted: Key Takeaway", [
    "\u2022 LangSmith: most complete self-hosted option (Enterprise plan)",
    "  \u2013 Full stack: frontend, backend, ClickHouse, PostgreSQL, Redis, blob storage",
    "  \u2013 Cloud guides for AWS, GCP, Azure",
    "",
    "\u2022 Braintrust: strong data residency but not fully self-hosted end-to-end",
    "  \u2013 Data plane runs in your infra; control plane remains managed",
    "",
    "\u2022 Arize: clear Phoenix self-hosting; AX self-hosting not clearly documented",
    "  \u2013 Phoenix: local, Docker, Kubernetes, or cloud",
])

# 9. Recommendations
section_slide("Scenario-Based Recommendations")

table_slide("Which Platform for Which Scenario?",
    ["Scenario", "Best Fit", "Why"],
    [
        ["Running agent (not LangGraph)", "Braintrust", "Best general fit without replatforming"],
        ["Debugging & quality improvement", "Arize AX", "Strongest improvement-loop framing"],
        ["LangGraph-based agent", "LangSmith", "Natural ecosystem alignment + runtime"],
        ["Review-heavy human workflow", "Braintrust", "Strongest review operations"],
        ["Agent platform + hosted runtime", "LangSmith", "Agent Server + deployment model"],
        ["Provider routing/caching needed", "Braintrust", "Gateway is clearest differentiator"],
        ["Standards-based telemetry priority", "Arize AX", "OTel + OpenInference + OTLP"],
    ]
)

# 10. Key Takeaways
section_slide("Key Takeaways")

content_slide("Summary", [
    "1. All three are strong and overlap significantly on core capabilities",
    "",
    "2. The choice depends on what you center your workflow around:",
    "   \u2022 Debugging & quality experiments \u2192 Arize AX",
    "   \u2022 End-to-end eval + review + inference ops \u2192 Braintrust",
    "   \u2022 Agent development + deployment + observability \u2192 LangSmith",
    "",
    "3. No wrong choice for observability/eval basics \u2014 all three deliver",
    "",
    "4. Differentiators that matter most:",
    "   \u2022 Arize: OTel standards, improvement loop, Alyx",
    "   \u2022 Braintrust: Gateway, review ops, classifiers, autoevals",
    "   \u2022 LangSmith: Agent Server, pairwise eval, deployment/runtime",
])

# Final slide
title_slide(
    "Thank You",
    "Based on vendor documentation review \u2022 June 2026"
)

# Save
output_path = "/Users/zhao/Desktop/Agent_observability/Agent_Observability_Platforms_Comparison.pptx"
prs.save(output_path)
print(f"Deck saved to: {output_path}")
